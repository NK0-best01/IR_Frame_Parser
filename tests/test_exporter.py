import datetime
import openpyxl
from irframe_parser.exporter import export_xlsx


def test_export_creates_valid_workbook(tmp_path):
    rows = [
        {
            "no": 1,
            "office": "สขจ.สกลนคร",
            "date": "26/08/2569",
            "status": "รออะไหล่",
            "sn": "ABC1234",
            "type": "Dell Optiplax 3050 AIO",
        },
        {
            "no": 2,
            "office": "พัทลุง",
            "date": "05/06/2553",
            "status": "รออะไหล่",
            "sn": "",
            "type": "Dell Optiplax 3050 AIO",
        },
    ]
    out_path = tmp_path / "export_test.xlsx"
    export_xlsx(rows, str(out_path))

    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    assert ws["A1"].value == "รายการเบิก IR Frame"
    assert ws.cell(row=5, column=2).value == "ลำดับ"
    assert ws.cell(row=6, column=3).value == "สขจ.สกลนคร"
    assert ws.cell(row=6, column=6).value == "ABC1234"


def test_missing_sn_footer_note_present(tmp_path):
    rows = [{"no": 1, "office": "A", "date": "", "status": "รออะไหล่", "sn": "", "type": "T"}]
    out_path = tmp_path / "export_missing_sn.xlsx"
    export_xlsx(rows, str(out_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    values = [c.value for row in ws.iter_rows() for c in row]
    assert any(v and "ไม่มีเลข SN" in str(v) for v in values)


def test_date_converted_to_datetime_cell(tmp_path):
    rows = [{"no": 1, "office": "A", "date": "26/08/2569", "status": "รออะไหล่", "sn": "SN1", "type": "T"}]
    out_path = tmp_path / "export_date.xlsx"
    export_xlsx(rows, str(out_path))
    wb = openpyxl.load_workbook(out_path)
    ws = wb.active
    date_cell = ws.cell(row=6, column=4)
    assert isinstance(date_cell.value, (datetime.date, datetime.datetime))
    assert date_cell.value.year == 2026  # 2569 (BE) - 543 = 2026 (CE)


def test_freeze_panes_set(tmp_path):
    rows = [{"no": 1, "office": "A", "date": "", "status": "รออะไหล่", "sn": "SN1", "type": "T"}]
    out_path = tmp_path / "export_freeze.xlsx"
    export_xlsx(rows, str(out_path))
    wb = openpyxl.load_workbook(out_path)
    assert wb.active.freeze_panes is not None
