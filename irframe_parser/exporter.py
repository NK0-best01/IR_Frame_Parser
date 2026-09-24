"""Excel exporter module for irframe_parser."""
from __future__ import annotations

import datetime
import re
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def _parse_thai_date(date_str: str) -> datetime.date | str:
    """Convert Thai BE date string (dd/mm/yyyy) into CE datetime.date object."""
    if not date_str:
        return ""
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", str(date_str).strip())
    if not m:
        return date_str
    day, month, year_val = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if year_val < 2400:
        year_val += 543
    gregorian_year = year_val - 543
    try:
        return datetime.date(gregorian_year, month, day)
    except ValueError:
        return date_str


def export_xlsx(rows: list[dict[str, Any]], out_path: str) -> None:
    """Export list of records into formatted Excel file matching the template."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ส่งเคสให้โม เพื่อเบิกจอทัช"

    # Set column widths
    col_widths = {
        "A": 3.5,
        "B": 8.0,
        "C": 20.0,
        "D": 15.0,
        "E": 18.0,
        "F": 22.0,
        "G": 28.0,
        "H": 3.5,
    }
    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    # Styles
    thin = Side(style="thin", color="777777")
    medium = Side(style="medium", color="555555")
    border_all_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
    border_all_medium = Border(left=medium, right=medium, top=medium, bottom=medium)

    fill_blue = PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid")
    fill_head = PatternFill(start_color="C9DAF8", end_color="C9DAF8", fill_type="solid")

    font_title = Font(name="AngsanaUPC", size=24, bold=True)
    font_sub = Font(name="Tahoma", size=12, bold=True, color="000000")
    font_header = Font(name="Angsana New", size=15, bold=True)
    font_data = Font(name="Angsana New", size=14)
    font_sig = Font(name="Tahoma", size=11)

    align_center = Alignment(horizontal="center", vertical="center")
    align_center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)

    # Title: A1:H2
    ws.merge_cells("A1:H2")
    ws["A1"] = "รายการเบิก IR Frame"
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    # Subtitle: B4:G4
    ws.row_dimensions[4].height = 27
    ws.merge_cells("B4:G4")
    subtitle_cell = ws["B4"]
    subtitle_cell.value = "รายการเบิกเฉพาะ IR frame (…..........................................)"
    subtitle_cell.font = font_sub
    subtitle_cell.fill = fill_blue
    subtitle_cell.alignment = align_center
    for c in range(2, 8):
        cell = ws.cell(row=4, column=c)
        cell.border = border_all_medium

    # Table Header at Row 5
    ws.row_dimensions[5].height = 40
    headers = [
        ("ลำดับ", 2),
        ("สนง.", 3),
        ("วันที่รับเคส", 4),
        ("สถานะ\nแจ้งซ่อม", 5),
        ("Serial Number", 6),
        ("เครื่อง/\nอุปกรณ์", 7),
    ]
    for text, col_idx in headers:
        cell = ws.cell(row=5, column=col_idx, value=text)
        cell.font = font_header
        cell.fill = fill_head
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_all_medium

    # Data Rows starting from row 6
    current_row = 6
    for idx, r in enumerate(rows, start=1):
        ws.row_dimensions[current_row].height = 22
        no_val = r.get("no", idx)
        office_val = r.get("office", "")
        date_obj = _parse_thai_date(r.get("date", ""))
        status_val = r.get("status", "")
        sn_val = r.get("sn", "")
        type_val = r.get("type", "")

        row_data = [
            (no_val, 2),
            (office_val, 3),
            (date_obj, 4),
            (status_val, 5),
            (sn_val, 6),
            (type_val, 7),
        ]

        for val, col_idx in row_data:
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_data
            cell.alignment = align_center_top
            cell.border = border_all_thin
            if isinstance(val, (datetime.date, datetime.datetime)):
                cell.number_format = "[$-1010000]d/m/yy;@"

        current_row += 1

    # Missing SN note if any
    if any(not r.get("sn") for r in rows):
        note_cell = ws.cell(row=current_row, column=2, value="***IR Frame ไม่มีเลข SN")
        note_cell.font = Font(name="Tahoma", size=12, bold=True, color="FF0000")
        note_cell.alignment = Alignment(horizontal="left", vertical="center")
        current_row += 1

    # Signature Block (Columns 5 to 7 to match table right boundary with complete enclosed border)
    sign_start_row = current_row + 1
    start_col = 5
    end_col = 7

    sig_lines = [
        ("ผู้ขอเบิกอุปกรณ์", sign_start_row),
        ("__________________", sign_start_row + 1),
        ("", sign_start_row + 2),
        ("_____/______/______", sign_start_row + 3),
    ]

    for text, r_idx in sig_lines:
        ws.row_dimensions[r_idx].height = 24
        ws.merge_cells(start_row=r_idx, start_column=start_col, end_row=r_idx, end_column=end_col)
        cell = ws.cell(row=r_idx, column=start_col, value=text)
        cell.font = font_sig
        cell.alignment = align_center

        for c_idx in range(start_col, end_col + 1):
            top_border = medium if r_idx == sign_start_row else None
            bottom_border = medium if r_idx == sign_start_row + 3 else None
            left_border = medium if c_idx == start_col else None
            right_border = medium if c_idx == end_col else None
            ws.cell(row=r_idx, column=c_idx).border = Border(
                top=top_border, bottom=bottom_border, left=left_border, right=right_border
            )

    # Freeze top 5 rows
    ws.freeze_panes = "A6"

    wb.save(out_path)
